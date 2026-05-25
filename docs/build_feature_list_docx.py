# -*- coding: utf-8 -*-
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = r"C:\imairuka\imairuka_share_20260507\docs\Imairuka_feature_list_20260526.docx"


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = "Yu Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_paragraph(doc, text, size=10.5, bold=False, color=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, size=9.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)


def add_table(doc, title, headers, rows, widths):
    add_paragraph(doc, title, size=14, bold=True, color=(0, 102, 204), space_after=8)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, header, bold=True, fill="EAF2FF")
        cell.width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            cells[i].width = Cm(widths[i])
    doc.add_paragraph()


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.6)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = doc.styles
styles["Normal"].font.name = "Yu Gothic"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Imairuka 機能一覧")
set_font(run, size=24, bold=True, color=(0, 86, 179))

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("ASP型ローンチ前確認版 / 2026年05月26日")
set_font(r, size=11, color=(90, 90, 90))

add_paragraph(
    doc,
    "本資料は、利用企業が Imairuka 案件管理システムで利用できる主要機能を整理したものです。案件、文書、決済、ユーザー、設定、外部連携まで、日常業務に必要な機能を利用企業向けにまとめています。",
    size=10.5,
    space_after=12,
)

add_table(
    doc,
    "1. 利用者向け 基本機能",
    ["カテゴリ", "機能", "概要"],
    [
        ["ダッシュボード", "利用状況の確認", "案件、文書、決済などの利用状況を確認します。"],
        ["案件一覧", "案件検索・一覧表示", "案件番号、顧客名、案件名、ステータス、タスク数、課題数、アサイン状況を一覧で確認します。"],
        ["案件詳細", "基本情報の確認", "顧客名、案件名、案件概要、ステータスなどを確認します。基本情報は折りたたみ可能です。"],
        ["プロジェクト管理", "ガント・タスク・課題・アサイン管理", "各案件に紐づく工程、詳細タスク、課題、担当者を管理します。"],
        ["AI下書き", "工程案の作成", "案件概要からAIで大まかなガントチャート案を作成します。会社ごとのAI APIキーを利用します。"],
    ],
    [3.2, 4.4, 10.2],
)

add_table(
    doc,
    "2. 文書・決済機能",
    ["カテゴリ", "機能", "概要"],
    [
        ["見積書", "作成・編集・プレビュー", "案件名を含めて見積書を作成し、PDFプレビューできます。"],
        ["納品書", "作成・管理", "案件に紐づく納品書を作成・管理します。"],
        ["請求書", "作成・管理", "案件に紐づく請求書を作成・管理します。"],
        ["領収書", "作成・管理", "入金後の領収書を作成・管理します。"],
        ["決済履歴", "Stripe Webhook保存", "Stripe決済成功イベントを受け取り、専用の決済履歴として保存・検索します。"],
        ["請求管理", "決済カテゴリ配下に配置", "請求関連の管理導線を決済管理カテゴリにまとめています。"],
    ],
    [3.2, 4.4, 10.2],
)

add_table(
    doc,
    "3. アカウント・権限",
    ["カテゴリ", "機能", "概要"],
    [
        ["複数アカウント", "ユーザー招待", "契約に応じて利用可能なユーザー数を表示し、招待できます。基本契約は1名まで、追加アカウントは別途課金想定です。"],
        ["権限", "オーナー・管理者・経理・一般・閲覧のみ", "ユーザーごとに権限を付与できます。閲覧者は編集・削除・追加操作が画面上でも非表示になり、サーバー側でも書き込みを拒否します。"],
        ["ユーザー枠", "契約中の利用枠表示", "標準ユーザー枠、追加ユーザー枠、招待中ユーザーを確認できます。"],
        ["招待管理", "招待URL発行・再発行", "権限を指定してユーザーを招待し、未承諾の招待を管理できます。"],
    ],
    [3.2, 4.4, 10.2],
)

add_table(
    doc,
    "4. 外部連携・セキュリティ",
    ["カテゴリ", "機能", "概要"],
    [
        ["Stripe", "Connect連携", "利用企業は自社のStripeアカウントを連携して、顧客からオンライン決済を受け付けられます。"],
        ["Stripe", "決済状態管理", "決済リンク作成、入金状態、返金済み決済の状態保持に対応します。"],
        ["AI設定", "OpenAI / Claude / Gemini", "会社ごとに利用するAIプロバイダーとAPIキーを設定します。Copilotは対象外です。"],
        ["会社単位のデータ管理", "自社データのみ表示", "利用企業ごとに案件・文書・決済・ユーザー情報を分離して管理します。"],
        ["セキュリティ", "閲覧者の書き込み制限", "閲覧権限ユーザーはサーバー側でも書き込み処理を拒否します。"],
        ["ログイン", "メールOTP対応", "OTPログイン有効化後は、メールアドレスとパスワードに加えて、登録メールアドレスへ届く認証コードを入力します。"],
        ["安定化", "Bootstrap背景レイヤー自動掃除", "ログイン後や画面遷移後に半透明の背景レイヤーが残らないよう、ページ初期化時に不要なbackdropを除去します。"],
    ],
    [3.2, 4.4, 10.2],
)

add_table(
    doc,
    "5. 利用開始前の確認事項",
    ["区分", "項目", "内容"],
    [
        ["初期設定", "会社情報", "社名、住所、連絡先、登録番号など、文書に表示する情報を確認します。"],
        ["初期設定", "ユーザー", "利用者、権限、招待中ユーザー、追加ユーザー枠を確認します。"],
        ["任意設定", "Stripe連携", "オンライン決済を利用する場合、自社のStripeアカウントを連携します。"],
        ["任意設定", "AI APIキー", "AI下書きを利用する場合、利用するAIプロバイダーとAPIキーを登録します。"],
        ["日常運用", "案件・文書", "顧客、商品、案件、見積書、納品書、請求書、領収書の作成フローを確認します。"],
    ],
    [2.4, 5.2, 10.2],
)

add_paragraph(
    doc,
    "備考: 本資料は機能整理用の一覧です。正式な販売資料では、画面キャプチャ、料金、サポート範囲、利用開始までの流れを追加します。",
    size=9.5,
    color=(90, 90, 90),
)

doc.save(OUT)
print(OUT)
