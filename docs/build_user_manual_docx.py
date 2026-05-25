from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
SCREENSHOT_DIR = BASE_DIR / "manual_screenshots"
OUTPUT = BASE_DIR / "Imairuka_user_manual_20260526.docx"
FONT = "Yu Gothic"


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, fill=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    set_font(run, 10, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    size = 18 if level == 1 else 13
    set_font(run, size, bold=True, color=(17, 24, 39))
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_font(run, 10.5, color=(31, 41, 55))
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(item)
        set_font(run, 10.2, color=(31, 41, 55))


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = 1
    cell = table.cell(0, 0)
    shade_cell(cell, "EEF6FF")
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(title)
    set_font(run, 10.5, bold=True, color=(0, 97, 255))
    paragraph.add_run("\n")
    run = paragraph.add_run(text)
    set_font(run, 9.6, color=(31, 41, 55))
    doc.add_paragraph()


def add_screenshot(doc, filename, caption):
    path = SCREENSHOT_DIR / filename
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.7))
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    set_font(caption_run, 8.5, color=(107, 114, 128))
    caption_paragraph.paragraph_format.space_after = Pt(8)


def add_section(doc, title, intro, bullets, screenshot, caption):
    add_heading(doc, title, 1)
    add_body(doc, intro)
    add_bullets(doc, bullets)
    add_screenshot(doc, screenshot, caption)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    for style_name in ["Normal", "List Bullet"]:
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(10.5)


def add_cover(doc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = 1
    cell = table.cell(0, 0)
    shade_cell(cell, "0D6EFD")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run("Imairuka 利用者向け操作マニュアル")
    set_font(run, 22, bold=True, color=(255, 255, 255))

    add_body(doc, "対象: Imairuka ASP/SaaS を利用する利用会社の管理者・担当者")
    add_body(doc, "更新日: 2026年5月26日")
    add_callout(
        doc,
        "このマニュアルについて",
        "実際の利用者画面のスクリーンショットを使い、ログインから案件、文書、請求、決済履歴、設定、ユーザー管理までの基本操作を確認できるようにしています。",
    )

    add_heading(doc, "全体の流れ", 1)
    rows = [
        ("1", "ログイン", "発行されたアカウントで利用者画面へ入ります。"),
        ("2", "案件登録・確認", "案件一覧で顧客、案件名、ステータスを確認します。"),
        ("3", "案件管理", "案件ごとにタスク、課題、アサイン、関連文書を管理します。"),
        ("4", "文書作成", "見積書、納品書、請求書、領収書を作成・確認します。"),
        ("5", "請求・決済確認", "請求管理と決済履歴で支払状況を確認します。"),
        ("6", "設定・ユーザー管理", "会社情報、Stripe連携、AI設定、ユーザー権限を管理します。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["順番", "機能", "内容"]
    for idx, header in enumerate(headers):
        set_cell_text(table.cell(0, idx), header, bold=True, fill="F3F4F6")
    for no, feature, detail in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], no)
        set_cell_text(cells[1], feature, bold=True)
        set_cell_text(cells[2], detail)


def build():
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    doc.add_page_break()
    add_section(
        doc,
        "1. ログイン",
        "発行されたメールアドレスとパスワードでログインします。ログイン後は、契約中の会社データだけを操作します。OTPログインが有効な場合は、登録メールアドレスへ届く認証コードも入力します。",
        [
            "ログイン情報は運営から案内されたものを使用します。",
            "利用者ごとの権限により、表示できるメニューや操作範囲が変わります。",
            "OTPが有効化された後は、メールアドレスとパスワード入力後に登録メールアドレスへ届く6桁コードを入力します。",
            "ログインできない場合は、メールアドレス、パスワード、アカウントの有効状態を確認します。",
        ],
        "01_login.png",
        "ログイン画面",
    )

    add_section(
        doc,
        "2. ダッシュボード",
        "ログイン後の最初の画面です。売上、未払い、案件状況、在庫アラートなどをまとめて確認できます。",
        [
            "当月売上や未払い請求書を確認し、優先対応が必要な項目を把握します。",
            "案件ステータスの分布を確認し、進行中・確認中・完了などの偏りを見ます。",
            "必要な操作は左メニューまたは上部メニューから移動します。",
        ],
        "02_dashboard.png",
        "ダッシュボード画面",
    )

    doc.add_page_break()
    add_section(
        doc,
        "3. 案件一覧",
        "案件一覧では、登録済みの案件を検索し、詳細確認や編集、削除を行います。",
        [
            "検索条件で顧客名、案件名、ステータスなどを絞り込みます。",
            "案件番号は現実的な形式で表示され、請求・文書・決済履歴との照合に使います。",
            "案件の詳細確認や編集は、行右側の操作ボタンから行います。",
        ],
        "03_orders_index.png",
        "案件一覧画面",
    )

    add_section(
        doc,
        "4. 案件管理",
        "案件管理では、各案件のプロジェクト進行に必要な情報をまとめて確認します。",
        [
            "管理対象案件一覧から、進行管理したい案件を開きます。",
            "案件概要、タスク、課題、アサイン、関連文書、決済状況を案件単位で確認します。",
            "案件概要からAIでガントチャートの下書きを作る場合は、設定画面でAI APIキーを登録しておきます。",
        ],
        "04_project_management_index.png",
        "案件管理の一覧画面",
    )

    add_screenshot(doc, "05_project_management_detail.png", "案件管理の詳細画面")

    doc.add_page_break()
    add_section(
        doc,
        "5. 請求管理",
        "請求管理では、案件ごとの請求金額、支払期限、支払方法、ステータスを確認できます。",
        [
            "入金確認や請求フォローの対象を一覧で把握します。",
            "銀行振込やカードなど、案件ごとの支払方法を確認します。",
            "詳細ボタンから個別の請求情報へ移動します。",
        ],
        "06_billing_orders.png",
        "請求管理画面",
    )

    add_section(
        doc,
        "6. 決済履歴",
        "決済履歴では、Stripe連携や入金情報に基づく決済状況を確認します。",
        [
            "決済済み、未入金などの状態で絞り込みできます。",
            "本番SaaSでは、Stripe Webhookで決済成功を受け取り、専用の決済履歴として保存する設計です。",
            "詳細ボタンから対象案件の決済情報を確認します。",
        ],
        "07_payment_histories.png",
        "決済履歴画面",
    )

    doc.add_page_break()
    add_section(
        doc,
        "7. 文書管理",
        "見積書、納品書、請求書、領収書を案件情報とひも付けて管理します。",
        [
            "各文書一覧から、詳細確認、編集、プレビュー、削除を行います。",
            "案件名の列で、どの案件に紐づく文書かを確認できます。",
            "必要に応じてPDF出力し、顧客提出用の文書として利用します。",
        ],
        "08_quotations.png",
        "見積書管理画面",
    )

    add_section(
        doc,
        "8. 設定",
        "設定画面では、会社情報、ステータス、Stripe連携、AI設定を管理します。",
        [
            "会社情報は各種文書や画面表示で使用します。",
            "Stripe連携は、利用会社が自社顧客からオンライン決済を受ける場合に使用します。",
            "AI設定では OpenAI、Claude、Gemini のAPIキーを会社ごとに登録できます。",
        ],
        "09_settings.png",
        "設定画面",
    )

    doc.add_page_break()
    add_section(
        doc,
        "9. ユーザー管理",
        "ユーザー管理では、契約中の利用枠、登録済みユーザー、未承認の招待を確認します。",
        [
            "基本契約は標準1名までです。追加アカウントは契約変更後に利用できます。",
            "ユーザーごとに権限を変更できます。不要になったアカウントは停止します。",
            "招待枠がない場合は、新規招待ボタンが使えない状態になります。",
        ],
        "10_account_users.png",
        "ユーザー管理画面",
    )

    add_heading(doc, "10. 運用上の注意", 1)
    add_bullets(
        doc,
        [
            "ImairukaのSaaS利用料は、運営またはベンダーからの請求書・銀行振込で管理します。",
            "Stripeは、利用会社が自社顧客からオンライン決済を受けるための機能です。",
            "追加アカウントはStripe課金ではなく、運営側で契約変更後に解放します。",
            "AI APIキーは各社ごとに保存されます。キーは画面上に再表示されません。",
            "OTPログインが有効な場合は、登録メールアドレスで認証コードを確認できる状態にしておきます。",
            "本マニュアルの画面は2026年5月26日時点のものです。実際の契約状態や権限により一部表示が異なる場合があります。",
        ],
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
